import io
import random
import re
from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

COLLEGE_DEFAULT = "Swami Vivekanand Institute of Engineering & Technology"
UMC_LINE = (
    "NOTE: Disclosure of Identity by writing Mobile No. or Marking of passing request "
    "on any paper of Answer Sheet will lead to UMC against the Student."
)
MARK_OPTIONS = [2, 4, 5, 8, 10]

BASE_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIR = BASE_DIR / "frontend"

app = FastAPI(title="SVIET Question Paper Studio", version="1.0.0")


# -------------------- Source question-bank logic --------------------
def normalize(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value)).strip()


def looks_like_heading(text: str) -> bool:
    value = normalize(text).strip(" :-")
    if not value:
        return True
    return bool(re.match(
        r"^(unit\s+[ivx\d]+|module\s+\d+|chapter\s+\d+|question\s*(bank|text)?|questions?|s\.?\s*no\.?|serial\s*no\.?)$",
        value, flags=re.IGNORECASE
    ))


def unique(items: list[str]) -> list[str]:
    result, seen = [], set()
    for item in items:
        text = normalize(item)
        if len(text) < 3 or looks_like_heading(text):
            continue
        key = text.casefold()
        if key not in seen:
            seen.add(key)
            result.append(text)
    return result


def read_docx(upload: io.BytesIO) -> list[str]:
    upload.seek(0)
    doc = Document(upload)
    items = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                items.append(max(cells, key=len))
    return unique(items)


def read_excel(upload: io.BytesIO) -> list[str]:
    upload.seek(0)
    workbook = pd.ExcelFile(upload)
    items = []
    for sheet_name in workbook.sheet_names:
        frame = pd.read_excel(workbook, sheet_name=sheet_name, header=None).fillna("")
        for _, row in frame.iterrows():
            cells = [normalize(v) for v in row.tolist() if normalize(v)]
            if cells:
                items.append(max(cells, key=len))
    return unique(items)


def read_bank(filename: str, raw: bytes) -> list[str]:
    upload = io.BytesIO(raw)
    name = filename.lower()
    if name.endswith(".docx"):
        return read_docx(upload)
    if name.endswith((".xlsx", ".xls")):
        return read_excel(upload)
    raise ValueError("Upload a Word (.docx) or Excel (.xlsx/.xls) file.")


# -------------------- Word document formatting logic --------------------
def add_paragraph(doc, text="", *, bold=False, size=10, align=None, space_after=0, italic=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(str(text))
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    return p


def number_to_words(n: int) -> str:
    words = {
        0: "ZERO", 1: "ONE", 2: "TWO", 3: "THREE", 4: "FOUR", 5: "FIVE",
        6: "SIX", 7: "SEVEN", 8: "EIGHT", 9: "NINE", 10: "TEN",
        11: "ELEVEN", 12: "TWELVE", 13: "THIRTEEN", 14: "FOURTEEN",
        15: "FIFTEEN", 16: "SIXTEEN", 17: "SEVENTEEN", 18: "EIGHTEEN",
        19: "NINETEEN", 20: "TWENTY"
    }
    return words.get(int(n), str(int(n)))


def add_centered_heading(doc, text):
    add_paragraph(doc, text, bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        from docx.oxml import OxmlElement
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def build_docx(meta, section_a, section_b, section_c, marks_b, marks_c, attempt_count):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)

    header_table = doc.add_table(rows=1, cols=2)
    header_table.autofit = False
    header_table.columns[0].width = Inches(3.5)
    header_table.columns[1].width = Inches(3.5)
    remove_table_borders(header_table)
    left_cell, right_cell = header_table.cell(0, 0), header_table.cell(0, 1)
    left_cell.width, right_cell.width = Inches(3.5), Inches(3.5)
    p_left = left_cell.paragraphs[0]
    p_left.paragraph_format.space_after = Pt(0)
    r_left = p_left.add_run("Roll No. __________________________________")
    r_left.font.name, r_left.font.size = "Times New Roman", Pt(10)
    p_right = right_cell.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_right.paragraph_format.space_after = Pt(0)
    r_right = p_right.add_run("No. of Pages: 1")
    r_right.font.name, r_right.font.size = "Times New Roman", Pt(10)

    add_paragraph(doc, f"Total No. of Questions: {1 + len(section_b) + len(section_c):02d}", size=10, space_after=0)
    add_paragraph(doc, meta["college"], bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_paragraph(doc, meta["degree_semester"], bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_paragraph(doc, meta["subject"], bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_paragraph(doc, f"Subject Code: {meta['subject_code']}", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    if meta["m_code"].strip():
        add_paragraph(doc, f"M. Code: {meta['m_code']}", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, f"Date of Examination: {meta['exam_date']}", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    time_table = doc.add_table(rows=1, cols=2)
    time_table.autofit = False
    time_table.columns[0].width = Inches(3.5)
    time_table.columns[1].width = Inches(3.5)
    remove_table_borders(time_table)
    left_cell, right_cell = time_table.cell(0, 0), time_table.cell(0, 1)
    left_cell.width, right_cell.width = Inches(3.5), Inches(3.5)
    p_left = left_cell.paragraphs[0]
    p_left.paragraph_format.space_after = Pt(0)
    r_left = p_left.add_run(f"Time: {meta['duration']}")
    r_left.font.name, r_left.font.size = "Times New Roman", Pt(10)
    p_right = right_cell.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_right.paragraph_format.space_after = Pt(0)
    r_right = p_right.add_run(f"Max. Marks: {meta['maximum_marks']}")
    r_right.font.name, r_right.font.size = "Times New Roman", Pt(10)

    add_paragraph(doc, "INSTRUCTIONS TO CANDIDATES:", bold=True, size=10, space_after=2)
    add_paragraph(doc, f"SECTION-A is COMPULSORY consisting of {number_to_words(len(section_a))} questions carrying TWO marks each.", size=10)
    add_paragraph(doc, f"SECTION-B has {number_to_words(len(section_b))} questions and SECTION-C has {number_to_words(len(section_c))} questions.", size=10)
    if marks_b == marks_c:
        add_paragraph(doc, f"Attempt any {number_to_words(min(attempt_count, len(section_b) + len(section_c)))} questions from SECTION B & C carrying {number_to_words(marks_b)} marks each.", size=10)
    else:
        add_paragraph(doc, f"SECTION-B questions carry {marks_b} marks each and SECTION-C questions carry {marks_c} marks each. Attempt questions as directed.", size=10)
    add_paragraph(doc, "Select at least TWO questions from SECTION-B & C.", size=10, space_after=5)

    add_centered_heading(doc, "SECTION-A")
    add_paragraph(doc, "1. Write short notes on:", size=10, space_after=3)
    sub_labels = "abcdefghijklmnopqrstuvwxyz"
    for index, question in enumerate(section_a):
        label = sub_labels[index] if index < len(sub_labels) else str(index + 1)
        add_paragraph(doc, f"({label}) {question}", size=10, space_after=5)

    add_centered_heading(doc, "SECTION-B")
    q_number = 2
    for question in section_b:
        add_paragraph(doc, f"{q_number}. {question}  [{marks_b} marks]", size=10, space_after=8)
        q_number += 1

    add_centered_heading(doc, "SECTION-C")
    for question in section_c:
        add_paragraph(doc, f"{q_number}. {question}  [{marks_c} marks]", size=10, space_after=8)
        q_number += 1

    # Keep the warning in the Word footer, near the bottom of the page.
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(UMC_LINE)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(8)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


# -------------------- API schemas and endpoints --------------------
class PaperRequest(BaseModel):
    meta: dict[str, Any]
    section_a: list[str] = Field(default_factory=list)
    section_b: list[str] = Field(default_factory=list)
    section_c: list[str] = Field(default_factory=list)
    marks_b: int = 8
    marks_c: int = 8
    attempt_count: int = 5


@app.get("/api/health")
def health():
    return {"status": "ok", "app": "SVIET Question Paper Studio"}


@app.get("/api/config")
def config():
    return {"mark_options": MARK_OPTIONS, "college_default": COLLEGE_DEFAULT, "umc_line": UMC_LINE}


@app.post("/api/parse-bank")
async def parse_bank(file: UploadFile = File(...)):
    try:
        raw = await file.read()
        questions = read_bank(file.filename or "upload", raw)
        if not questions:
            raise HTTPException(status_code=422, detail="No questions were detected in this file. Try a Word or Excel file with one question per paragraph/row.")
        return {"filename": file.filename, "count": len(questions), "questions": questions}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read this question bank: {exc}") from exc


@app.post("/api/generate-paper")
def generate_paper(request: PaperRequest):
    if request.marks_b not in MARK_OPTIONS or request.marks_c not in MARK_OPTIONS:
        raise HTTPException(status_code=422, detail=f"Marks must be one of {MARK_OPTIONS}.")
    meta = request.meta
    required = ["college", "degree_semester", "subject", "subject_code", "m_code", "exam_date", "duration", "maximum_marks"]
    missing = [key for key in required if key not in meta]
    if missing:
        raise HTTPException(status_code=422, detail=f"Missing paper details: {', '.join(missing)}")
    a, b, c = [unique(list(map(str, qs))) for qs in (request.section_a, request.section_b, request.section_c)]
    if not (a or b or c):
        raise HTTPException(status_code=422, detail="Select at least one question before downloading.")
    if set(b).intersection(c):
        raise HTTPException(status_code=422, detail="The same long question cannot appear in both Section B and Section C.")
    try:
        content = build_docx(meta, a, b, c, request.marks_b, request.marks_c, max(0, request.attempt_count))
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Could not create the Word paper: {exc}") from exc
    safe_subject = re.sub(r"[^A-Za-z0-9_-]+", "_", str(meta.get("subject", "SVIET_Exam"))).strip("_") or "SVIET_Exam"
    headers = {"Content-Disposition": f'attachment; filename="{safe_subject}_Question_Paper.docx"'}
    return StreamingResponse(io.BytesIO(content), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers=headers)


@app.get("/")
def index():
    return FileResponse(FRONTEND_DIR / "index.html")


app.mount("/assets", StaticFiles(directory=FRONTEND_DIR), name="assets")
